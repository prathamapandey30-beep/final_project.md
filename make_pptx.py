from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ─────────────────────────────────────────────────────────────────
GREEN      = RGBColor(0x2E, 0x7D, 0x32)
GREEN_LITE = RGBColor(0x4C, 0xAF, 0x50)
GREEN_BG   = RGBColor(0xE8, 0xF5, 0xE9)
GREEN_PALE = RGBColor(0xF0, 0xF7, 0xF0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK       = RGBColor(0x1A, 0x2B, 0x1B)
GRAY       = RGBColor(0x55, 0x77, 0x57)
AMBER      = RGBColor(0xF5, 0x9E, 0x0B)
ROW_ALT    = RGBColor(0xF4, 0xFA, 0xF4)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ── Helpers ──────────────────────────────────────────────────────────────────

def rect(sl, l, t, w, h, fill=None, border=None, border_w=1.5):
    s = sl.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(border_w)
    else:
        s.line.fill.background()
    return s

def txt(sl, text, l, t, w, h,
        size=14, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, italic=False):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run(); r.text = text
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.color.rgb = color
    r.font.italic = italic
    r.font.name   = "Segoe UI"
    return tb

def header(sl, title, subtitle=""):
    """Standard green top bar."""
    rect(sl, 0, 0, 13.33, 1.25, fill=GREEN)
    txt(sl, title,    0.4, 0.08, 10, 0.7,  size=30, bold=True, color=WHITE)
    if subtitle:
        txt(sl, subtitle, 0.4, 0.8, 11, 0.38, size=12, color=GREEN_BG, italic=True)

def card(sl, l, t, w, h, title, bullets,
         title_color=GREEN, bg=GREEN_BG, border=GREEN_LITE,
         title_size=13, bullet_size=11):
    """A simple labelled card with bullet lines."""
    rect(sl, l, t, w, h, fill=bg, border=border)
    txt(sl, title, l+0.18, t+0.12, w-0.28, 0.42, size=title_size, bold=True, color=title_color)
    body = "\n".join(bullets)
    txt(sl, body, l+0.18, t+0.58, w-0.28, h-0.68, size=bullet_size, color=DARK)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
rect(sl, 0, 0, 4.6,   7.5, fill=GREEN)          # left green panel

txt(sl, "GreenPlate",            0.15, 2.1, 4.3, 0.9,  size=38, bold=True,  color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Food Waste\nReduction App", 0.15, 3.05, 4.3, 0.9, size=17, color=GREEN_BG, align=PP_ALIGN.CENTER, italic=True)

rect(sl, 5.0, 2.75, 7.9, 0.05, fill=GREEN_LITE)  # accent line

txt(sl, "Scan receipts.\nTrack expiry.\nReduce waste.",
    5.0, 1.5, 7.9, 1.5, size=24, bold=True, color=GREEN)

txt(sl, "A mobile-first app that helps Finnish households\nautomatically track groceries, estimate shelf life,\nand act before food expires.",
    5.0, 3.05, 7.9, 1.6, size=13, color=DARK)

txt(sl, "AI in Practice  |  Xamk  |  2026",
    5.0, 6.9, 7.9, 0.4, size=11, color=GRAY)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — THE PROBLEM
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
header(sl, "The Problem", "Why food waste is a household crisis that needs a smarter tool")

problems = [
    ("Forgotten Food",
     ["Items expire unseen at the back of shelves", "No visibility into what you own"]),
    ("Manual Tracking is Tedious",
     ["Existing apps require typing every product", "Users abandon them within days"]),
    ("Unknown Shelf Life",
     ["Consumers don't know how long food lasts", "This leads to early disposal of safe food"]),
    ("No Actionable Next Step",
     ["Even when aware, people don't know what to cook", "Expiring ingredients go to waste anyway"]),
]

# 2x2 grid — 4 cards
CARD_W = 6.2; CARD_H = 2.5
for i, (title, bullets) in enumerate(problems):
    col = i % 2;  row = i // 2
    x = 0.3  + col * 6.5
    y = 1.35 + row * 2.75
    card(sl, x, y, CARD_W, CARD_H, title, bullets,
         title_size=14, bullet_size=12)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — THE SOLUTION
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
header(sl, "The Solution", "GreenPlate automates the entire food management cycle")

# 5-step flow
steps = ["Scan Receipt", "Extract Items", "Update Pantry", "Expiry Alerts", "Recipe Ideas"]
for i, label in enumerate(steps):
    x = 0.3 + i * 2.55
    rect(sl, x, 1.4, 2.2, 1.7, fill=GREEN_BG, border=GREEN_LITE)
    txt(sl, str(i+1), x+0.9, 1.5, 0.4, 0.4, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rect(sl, x+0.82, 1.5, 0.42, 0.38, fill=GREEN_LITE)
    txt(sl, label, x+0.1, 1.98, 2.0, 0.8, size=12, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    if i < 4:
        txt(sl, "->", x+2.2, 2.0, 0.38, 0.5, size=18, bold=True, color=GREEN_LITE)

# 3 key principles
principles = [
    ("Privacy-First",
     "All OCR runs on-device. No images or data\nleave your phone. Fully GDPR compliant."),
    ("Zero Manual Entry",
     "Receipts are scanned automatically. No typing\nor manual data entry ever required."),
    ("Actionable Alerts",
     "Get notified before food expires, with a\npersonalised recipe suggestion ready."),
]
for i, (title, desc) in enumerate(principles):
    x = 0.3 + i * 4.35
    rect(sl, x, 3.4, 4.1, 3.75, fill=GREEN_PALE, border=GREEN)
    txt(sl, title, x+0.2, 3.52, 3.7, 0.48, size=14, bold=True, color=GREEN)
    txt(sl, desc,  x+0.2, 4.08, 3.7, 2.9,  size=12, color=DARK)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — CORE FEATURES  (fixed 3-column × 2-row grid)
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
header(sl, "Core Features", "Every feature built into the MVP")

features = [
    ("Receipt Scanner",
     ["On-device OCR — Tesseract.js", "No cloud, no privacy risk", "Review & edit before saving"]),
    ("Smart Pantry",
     ["Auto-categorised inventory", "Colour-coded expiry bars", "Search, filter & sort items"]),
    ("Expiry Alerts",
     ["Red / amber urgency badges", "Dashboard 'Use First' section", "Expired item detection"]),
    ("Smart Recipes",
     ["Ranked by expiry urgency", "Match % shown per recipe", "Step-by-step cooking guide"]),
    ("Impact Dashboard",
     ["Money saved (EUR)", "Food waste reduced (kg)", "Monthly goal progress ring"]),
    ("Gamification",
     ["Daily eco-streak counter", "Achievement badges", "4-slide onboarding flow"]),
]

# Fixed grid: 3 columns, 2 rows
COLS   = 3
CARD_W = 4.08
CARD_H = 2.55
GAP_X  = 0.24   # horizontal gap
GAP_Y  = 0.20   # vertical gap
START_X = 0.25
START_Y = 1.38

for i, (title, bullets) in enumerate(features):
    col = i % COLS
    row = i // COLS              # ← FIXED (was i // 2)
    x = START_X + col * (CARD_W + GAP_X)
    y = START_Y + row * (CARD_H + GAP_Y)
    card(sl, x, y, CARD_W, CARD_H, title, bullets,
         title_size=13, bullet_size=11)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — TECH STACK
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
header(sl, "Technology Stack", "Lean, fast, and built for privacy")

rows = [
    ("Frontend",  "React 18 + Vite",   "Fast SPA with HMR development server"),
    ("Routing",   "React Router v6",   "Client-side navigation between pages"),
    ("OCR",       "Tesseract.js",      "Fully on-device text recognition — no cloud"),
    ("Styling",   "Vanilla CSS",       "Custom design system with CSS variables"),
    ("Icons",     "Lucide React",      "Clean, consistent icon library"),
    ("Dates",     "date-fns",          "Lightweight expiry date calculations"),
    ("Storage",   "localStorage",      "Persistent state — GDPR-friendly, no backend"),
    ("IDs",       "UUID v4",           "Collision-free IDs for every inventory item"),
]

col_labels = ["Layer", "Tool / Library", "Purpose"]
col_x      = [0.3, 2.8, 6.2]
col_w      = [2.3, 3.2, 6.8]
ROW_H      = 0.62
START_Y    = 1.38

# Header row
rect(sl, 0.3, START_Y, 12.73, ROW_H, fill=GREEN)
for j, label in enumerate(col_labels):
    txt(sl, label, col_x[j], START_Y+0.1, col_w[j], 0.42,
        size=13, bold=True, color=WHITE)

# Data rows
for i, (layer, tool, purpose) in enumerate(rows):
    y    = START_Y + (i+1) * ROW_H
    fill = GREEN_BG if i % 2 == 0 else WHITE
    rect(sl, 0.3, y, 12.73, ROW_H, fill=fill, border=GREEN_BG, border_w=0.5)
    txt(sl, layer,   col_x[0], y+0.1, col_w[0], 0.42, size=12, bold=True,  color=GREEN)
    txt(sl, tool,    col_x[1], y+0.1, col_w[1], 0.42, size=12, bold=False, color=DARK)
    txt(sl, purpose, col_x[2], y+0.1, col_w[2], 0.42, size=12, bold=False, color=GRAY)

# Future footer
rect(sl, 0.3, 7.08, 12.73, 0.32, fill=RGBColor(0xFF, 0xF8, 0xE1))
txt(sl, "Future: Firebase Firestore  |  FCM Push Notifications  |  Google ML Kit  |  Gemini AI",
    0.5, 7.09, 12.5, 0.3, size=10, color=RGBColor(0x78, 0x35, 0x0F), italic=True)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — TARGET AUDIENCE
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, 13.33, 7.5, fill=WHITE)
header(sl, "Target Audience", "Who GreenPlate is built for")

segments = [
    ("Finnish Households",
     GREEN,
     ["Primary target market", "Average EUR 500+ wasted per year", "Need automation, not manual effort", "Motivated by cost savings"]),
    ("Digitally-Savvy Adults",
     RGBColor(0x15, 0x65, 0xC0),
     ["Age 25-45, smartphone-native", "Comfortable adopting new apps", "Actively seek sustainable tools", "Respond well to gamification"]),
    ("Eco-Conscious Users",
     RGBColor(0x6A, 0x1B, 0x9A),
     ["Feel guilty about food waste", "Want to measure their impact", "Motivated by CO2 & waste stats", "Value privacy and transparency"]),
]

SEG_W = 4.0; SEG_H = 4.7
for i, (title, color, bullets) in enumerate(segments):
    x = 0.35 + i * 4.35
    rect(sl, x, 1.35, SEG_W, SEG_H, fill=GREEN_PALE, border=color)
    rect(sl, x, 1.35, SEG_W, 0.65, fill=color)   # coloured title bar
    txt(sl, title, x+0.15, 1.43, SEG_W-0.2, 0.5,
        size=14, bold=True, color=WHITE)
    for j, b in enumerate(bullets):
        txt(sl, "  " + b, x+0.15, 2.18+j*0.72, SEG_W-0.25, 0.62,
            size=12, color=DARK)

# B2B note at bottom
rect(sl, 0.35, 6.25, 12.6, 1.0, fill=GREEN_BG, border=GREEN_LITE)
txt(sl, "B2B Opportunities (Future Roadmap)",
    0.55, 6.3, 8, 0.4, size=13, bold=True, color=GREEN)
txt(sl, "Grocery retailers (affiliate promotions)  |  Employers & canteens (bulk licensing)  |  Municipal sustainability programs",
    0.55, 6.72, 12.3, 0.45, size=11, color=DARK)

# ── Save ─────────────────────────────────────────────────────────────────────
out = r"c:\Users\prath\Documents\Xamk Courses\AI in Practice\Lessons\Final Project App\GreenPlate_6Slides.pptx"
prs.save(out)
print("Saved: " + out)
